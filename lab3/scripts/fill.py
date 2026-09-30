import copy, docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from docx.text.paragraph import Paragraph
SRC='/root/.claude/uploads/de2f076b-18ee-50d2-8724-46a3b7c13fa0/bda20359-Lab_3_report.docx'
d=docx.Document(SRC)
BLUE=None
def font(run, code=False, bold=False, color=None, size=None):
    rpr=run._r.get_or_add_rPr(); f=rpr.find(qn('w:rFonts'))
    if f is None: f=OxmlElement('w:rFonts'); rpr.insert(0,f)
    a='Courier New' if code else 'Times New Roman'
    for k,v in [('w:ascii',a),('w:hAnsi',a),('w:cs',a),('w:eastAsia','標楷體')]: f.set(qn(k),v)
    run.bold=bold
    if color: run.font.color.rgb=color
    if size: run.font.size=Pt(size)
def new_par_after(p):
    np=OxmlElement('w:p'); p._p.addnext(np); return Paragraph(np,p._parent)
class Writer:
    def __init__(s,anchor): s.cur=anchor
    def par(s, parts, indent=0, align=None, space_after=4):
        p=new_par_after(s.cur); s.cur=p
        if indent: p.paragraph_format.left_indent=Inches(indent)
        p.paragraph_format.space_after=Pt(space_after)
        if align: p.alignment=align
        if isinstance(parts,str): parts=[parts]
        for part in parts:
            if isinstance(part,str): part=(part,{})
            t,o=part; r=p.add_run(t); font(r,**o)
        return p
    def img(s,path,w=5.75,caption=None):
        p=new_par_after(s.cur); s.cur=p; p.alignment=WD_ALIGN_PARAGRAPH.CENTER
        p.add_run().add_picture(path,width=Inches(w))
        if caption: s.par([(caption,{'size':10})],align=WD_ALIGN_PARAGRAPH.CENTER,space_after=8)
C={'code':True}; B={'bold':True}
body=d.element.body
def paras(): return d.paragraphs
def find(text):
    for p in paras():
        if p.text.strip().startswith(text): return p
    raise KeyError(text)
def clear_blanks_after(p):
    # remove consecutive empty paragraphs following p (keep none)
    nxt=p._p.getnext()
    while nxt is not None and nxt.tag==qn('w:p') and not ''.join(t.text or '' for t in nxt.iter(qn('w:t'))).strip() and not list(nxt.iter(qn('w:drawing'))):
        n2=nxt.getnext(); body.remove(nxt); nxt=n2
def blank(w): w.par('',space_after=0)

# ---------- header
_ts=list(d.paragraphs[0]._p.iter(qn('w:t')))
for a_,b_ in zip(_ts,_ts[1:]):
    if a_.text=='學號' and b_.text==':_____': b_.text=': 112652006'
    if a_.text=='姓名' and b_.text==':_____': b_.text=': 陳昱翰'
P=lambda w,t,**k: w.par(t,**k)

# ---------- Q1
a=find('當訂閱的topic為 /a/test'); clear_blanks_after(a); w=Writer(a)
P(w,'Ans: 不能。')
P(w,'因為topic是用 "/" 來分層的，/a/test 開頭多了一個 "/"，所以第一層其實是空的，總共有三層（空、a、test），而 a/test 只有兩層（a、test）。沒有用萬用字元的話，每一層都要一樣才會收到，所以訂閱 /a/test 收不到 a/test 的訊息。')
P(w,'實際測試時，pub 到 a/test 的 test_q1，訂閱 /a/test 的那邊沒有收到。')
w.img('img_q1_vm.png',w=4.5,caption='圖1  Q1測試結果')
blank(w)

# ---------- Q2
a=find('所有roof的所有內容'); clear_blanks_after(a); w=Writer(a)
P(w,'Ans:')
for i,t in enumerate(['+/roof/brightness/day','house1/firstfloor/#','house2/+/+/night','+/roof/#'],1):
    P(w,f'{i}. {t}',indent=0.2,space_after=0)
blank(w)
P(w,'+ 只能代表一層，# 可以代表後面所有層，但只能放在最後面。第1、3題是中間某幾層不限定，所以用 +；第2、4題是要某一層以下的全部內容，所以用 #。')
P(w,'我把36種topic全部pub一次，四個subscriber分別收到2、6、9、12筆，跟算出來的一樣。')
w.img('img_q2_vm.png',w=4.3,caption='圖2  Q2測試結果')
blank(w)

# ---------- Q3
a=find('排除匿名使用者截圖'); clear_blanks_after(a); w=Writer(a)
P(w,'先用 mosquitto_passwd 建立帳號（帳號112652006，密碼123456），再到 mosquitto.conf 加上 password_file /etc/mosquitto/passwd 和 allow_anonymous false，然後 restart mosquitto。')
P(w,'之後不加帳密直接 sub 或 pub 都會出現 not authorised，密碼打錯也一樣，加上 -u 112652006 -P 123456 才能正常收發。')
w.img('img_q3a_vm.png',w=5.2,caption='圖3  建立帳號')
w.img('img_q3b_vm.png',w=4.6,caption='圖4  修改設定檔並重啟')
w.img('img_q3c_vm.png',w=4.6,caption='圖5  匿名被拒絕，用帳密可以正常收發')
blank(w)

# ---------- Q4 table
tbl=d.tables[0]
res=[('in  msg1','無'),('out  msg2','fromA/out  msg2'),('out/msg  msg3','fromA/msg  msg3'),
     ('/out  msg4','無'),('fromB/out  msg5','out  msg5'),('fromB/fromB  msg6','fromB  msg6')]
for row,(ra,rb) in zip(tbl.rows[1:],res):
    for cell,val in zip(row.cells[1:],(ra,rb)):
        r=cell.paragraphs[0].add_run(val); font(r,size=11)
anchor=OxmlElement('w:p'); tbl._tbl.addnext(anchor); w=Writer(Paragraph(anchor,d._body))
P(w,'結果討論：')
for t in ['1. A的subscriber訂閱 #，所以A自己pub的訊息都會照原本的topic顯示。',
          '2. topic # out 1 out/ fromA/ 的意思是A上 out/ 開頭的訊息會送到B，而且 out/ 會換成 fromA/，所以msg3在B變成 fromA/msg。',
          '3. msg2的topic只有 out，因為 # 也會匹配上一層本身，所以 out 也符合 out/#，會被送到B，只是它沒有 out/ 可以換，就直接在前面加上 fromA/，B收到的是 fromA/out。',
          '4. msg1（in）和msg4（/out）都不符合 out/#，所以只有A收得到，B收不到。',
          '5. topic # in 1 fromB/ 只有設local-prefix，所以B上的所有訊息都會傳回A，而且前面加上 fromB/，msg5、msg6在A分別變成 fromB/out 和 fromB/fromB，B自己則是照原本的topic顯示。']:
    P(w,t,indent=0.2)
w.img('img_q4_vm.png',w=5.7,caption='圖6  Q4測試結果')
clear_blanks_after(w.cur); blank(w)
body.remove(anchor)

# ---------- Q5
a=find('在MQTT bridge configuration中，QoS level'); clear_blanks_after(a); w=Writer(a)
P(w,'Ans: QoS有0、1、2三種。')
for t in ['QoS 0：最多送一次。送出去就不管了，對方也不用回應，所以有可能掉訊息，但速度最快。',
          'QoS 1：至少送一次。對方收到要回PUBACK，沒收到回應就會重送，所以訊息不會掉，但可能會收到重複的。',
          'QoS 2：剛好送一次。要經過PUBLISH、PUBREC、PUBREL、PUBCOMP四次交握，確保不會掉也不會重複，但最慢、負擔也最大。']:
    P(w,t,indent=0.2)
P(w,'在bridge設定裡，topic後面的那個數字就是兩個broker之間傳這些topic時用的QoS，這次實驗設的是1。')
blank(w)

# ---------- Q6
a=find('什麼是"保留消息'); clear_blanks_after(a); w=Writer(a)
P(w,'Ans:')
P(w,'1. 設成true的話，bridge斷線時遠端broker會把它的訂閱和還沒送出的訊息都清掉，重新連線後要重新訂閱，斷線期間的訊息就收不到了。設成false（預設）的話，遠端broker會保留bridge的訂閱，斷線期間QoS 1、2的訊息也會先存起來，重連之後再補送。')
P(w,'2. 如果是false，改了bridge的topic之後重連，還是會收到舊topic的訊息。因為訂閱是存在遠端broker上的，重連時會沿用舊的session，而bridge只會去訂閱新的topic，不會取消舊的，所以舊的訂閱還在。解決方法就是先把cleansession改成true重連一次，把舊的清掉，再改回false。')
P(w,'3. retained message是publish時有加retain旗標（例如 mosquitto_pub -r）的訊息，broker會把每個topic最後一筆retained message存起來，之後只要有人訂閱這個topic就會馬上收到這一筆。設成true的話每次重連都要重新訂閱，而每訂閱一次，broker就會把符合的retained message全部再送一次。bridge通常是訂閱 # 這種範圍很大的topic，所以只要重連次數一多，就會一直收到一大堆retained message。')
blank(w)

# ---------- Q7
a=find('Q7.心得'); w=Writer(a)
P(w,'這次實驗學到MQTT的基本用法。一開始在設定VM的時候，因為VM沒有關機，網路跟處理器的設定都不能改，關機後才改得了。做pub/sub的時候，我一開始把sub和pub開在同一個terminal，sub跑起來之後按Ctrl+C跳出來再pub，當然收不到，後來才知道要開兩個terminal，一個負責訂閱、一個負責發布。')
P(w,'Q1原本覺得 /a/test 跟 a/test 應該差不多，實際測了才知道開頭多一個 / 就會多一層。Q2用 + 和 # 去篩選topic還蠻直觀的。bridge的部分比較複雜，out、in再加上prefix的轉換要想一下，像msg2的topic是 out 也會被送出去這點一開始沒想到，後來才搞懂。整體來說MQTT設定起來不難，pub跟sub只要知道broker的IP就可以溝通，蠻適合用在IoT上的。')

d.save('Lab_3_report_filled.docx'); print('saved')
